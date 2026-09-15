#!/usr/bin/env python3
"""Generate enigmatic plot + TURBO TEMI 360 Word document from cipher-data.json."""
from __future__ import annotations

import json
import math
from pathlib import Path

from PIL import Image, ImageDraw, ImageFilter, ImageFont
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path("/workspace")
PUB = ROOT / "public"
DL = PUB / "downloads"
DL.mkdir(parents=True, exist_ok=True)
data = json.loads((ROOT / "src/lib/cipher-data.json").read_text(encoding="utf-8"))

Pmod = data["curve"]["p"]
A = data["curve"]["a"]
B = data["curve"]["b"]
ORDER = data["curve"]["order"]
G = tuple(data["curve"]["G"])
encoded = data["messages"]
multiples = data["multiples"]
day_enc = data["days"]
st = data["stats"]
points = [tuple(p) for p in data["points"]]


def y_upper(x: float) -> float:
    v = x**3 + 2 * x + 3
    return math.sqrt(v) if v >= 0 else 0.0


# ---------- enigmatic plot ----------
W, H = 1600, 1000
img = Image.new("RGB", (W, H), (12, 14, 16))
draw = ImageDraw.Draw(img)
for gy in range(0, H, 80):
    draw.line([(0, gy), (W, gy)], fill=(28, 32, 36), width=1)
for gx in range(0, W, 80):
    draw.line([(gx, 0), (gx, H)], fill=(28, 32, 36), width=1)


def mapx(x: float) -> int:
    return int(180 + (x + 1.2) / 7.4 * (W - 320))


def mapy(y: float) -> int:
    return int(H / 2 - y * 38)


draw.line([(mapx(-1.2), mapy(0)), (mapx(6.1), mapy(0))], fill=(90, 96, 100), width=1)
draw.line([(mapx(0), mapy(-12)), (mapx(0), mapy(12))], fill=(90, 96, 100), width=1)

curve_u, curve_l = [], []
x = -1.0
while x <= 6.05:
    yu = y_upper(x)
    curve_u.append((mapx(x), mapy(yu)))
    curve_l.append((mapx(x), mapy(-yu)))
    x += 0.02
draw.line(curve_u, fill=(180, 196, 190), width=3)
draw.line(curve_l, fill=(110, 126, 122), width=2)

star_pts = []
for m in encoded:
    pt = m["point"]
    if pt is None:
        sx, sy = mapx(2.4), 70
        star_pts.append((sx, sy, m, True))
        for r, col in [(16, (50, 42, 40)), (8, (160, 130, 120)), (3, (236, 231, 222))]:
            draw.ellipse((sx - r, sy - r, sx + r, sy + r), fill=col)
        continue
    xr = 0.15 + pt[0] / 37 * 5.4
    yu = y_upper(xr)
    if pt[1] > Pmod / 2:
        yu = -yu
    sx, sy = mapx(xr), mapy(yu)
    star_pts.append((sx, sy, m, False))
    for r, col in [(18, (40, 55, 52)), (10, (90, 130, 120)), (4, (230, 232, 220))]:
        draw.ellipse((sx - r, sy - r, sx + r, sy + r), fill=col)

path = [(p[0], p[1]) for p in star_pts]
if len(path) >= 2:
    draw.line(path, fill=(90, 110, 104), width=1)

try:
    font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 14)
    font_sm = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 12)
    font_lg = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 22)
except OSError:
    font = font_sm = font_lg = ImageFont.load_default()

for sx, sy, m, inf in star_pts:
    label = f"{m['id']}" + (" ∞" if inf else "")
    draw.text((sx + 10, sy - 16), label, fill=(236, 231, 222), font=font)

draw.rectangle((48, 36, 820, 122), fill=(18, 20, 22))
draw.text((64, 48), "CONSTELACAO SEM JAULA", fill=(236, 231, 222), font=font_lg)
sdx = st["sd_x"]
integ = st["integral_-1_to_4"]
legend = f"y^2 = x^3 + 2x + 3   ·   18 mensagens -> kG   ·   sigma_x = {sdx:.3f}"
draw.text((64, 84), legend, fill=(155, 149, 138), font=font_sm)
foot1 = (
    f"Soma das gematrias = {st['sum_of_key_gematria']}   ·   "
    f"Soma dos k = {st['sum_of_k']}   ·   "
    f"integral [-1,4] = {integ:.4f}"
)
draw.text((64, H - 70), foot1, fill=(155, 149, 138), font=font_sm)
draw.text(
    (64, H - 46),
    "Decifracao: k = (soma A=1..Z + indice) mod 20, ponto = k*(4,1) em F37",
    fill=(111, 158, 152),
    font=font_sm,
)

img = img.filter(ImageFilter.SMOOTH)
out_img = PUB / "cifra-enigmatica.png"
img.save(out_img, "PNG", optimize=True)
img_doc_path = DL / "cifra-enigmatica-doc.png"
img.resize((1200, 750), Image.Resampling.LANCZOS).save(img_doc_path, "PNG", optimize=True)
print("image", out_img, out_img.stat().st_size)

# ---------- Word ----------
doc = Document()
section = doc.sections[0]
section.page_width = Cm(21.0)
section.page_height = Cm(29.7)
section.left_margin = Cm(1.8)
section.right_margin = Cm(1.8)
section.top_margin = Cm(1.6)
section.bottom_margin = Cm(1.6)

INK = (12, 14, 16)
SAGE = (70, 110, 104)
MUTED = (90, 86, 80)


def set_run(run, size=11, bold=False, color=INK, italic=False, name="Calibri"):
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.name = name
    run.font.color.rgb = RGBColor(*color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)


def shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "C8C4BC")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def h(text, level=1):
    p = doc.add_heading(text, level=level)
    for r in p.runs:
        set_run(r, size=16 if level == 1 else 13, bold=True, color=INK if level == 1 else SAGE)
    return p


def para(text, size=11, bold=False, italic=False, color=INK, space_after=6):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run(r, size=size, bold=bold, italic=italic, color=color)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    return p


def kv_table(rows, headers=None):
    cols = len(rows[0])
    t = doc.add_table(rows=(1 + len(rows) if headers else len(rows)), cols=cols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    start = 0
    if headers:
        for j, htext in enumerate(headers):
            cell = t.rows[0].cells[j]
            cell.text = ""
            r = cell.paragraphs[0].add_run(htext)
            set_run(r, size=9, bold=True, color=(255, 255, 255))
            shade_cell(cell, "0C0E10")
            set_cell_border(cell)
        start = 1
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.rows[start + i].cells[j]
            cell.text = ""
            r = cell.paragraphs[0].add_run(str(val))
            set_run(r, size=9, bold=(j == 0), color=INK)
            shade_cell(cell, "F4F1EA" if i % 2 else "FFFFFF")
            set_cell_border(cell)
    doc.add_paragraph()
    return t


p = doc.add_paragraph()
r = p.add_run("TURBO TEMI 360  ·  PROTOCOLO EDITORIAL")
set_run(r, size=9, bold=True, color=SAGE)

p = doc.add_paragraph()
r = p.add_run("BÍBLIA SEM JAULA 365")
set_run(r, size=26, bold=True, color=INK)
p.paragraph_format.space_after = Pt(0)

p = doc.add_paragraph()
r = p.add_run("Cifra elíptica pedagógica  ·  Curva Almeida  y² ≡ x³ + 2x + 3 (mod 37)")
set_run(r, size=12, italic=True, color=SAGE)

para(
    "Documento simples, organizado e calculável à mão. Não é criptografia de produção. "
    "Serve para ensinar o formato elíptico, o logaritmo discreto e o vínculo entre a leitura "
    "bíblica e um ponto no grupo da curva.",
    size=11,
    color=MUTED,
)

kv_table(
    [
        ["ID", "CIFRA-365-JFA-ECC-TOY"],
        ["Fonte", "Bíblia Sem Jaula 365 — João Ferreira de Almeida"],
        ["Método", "ECC pedagógica (Koblitz/Miller, 1985) — Wikipédia: Criptografia de curva elíptica"],
        ["Corpo", "F37  (primo p = 37)"],
        ["Curva", "E: y² = x³ + 2x + 3"],
        ["Discriminante", "4a³ + 27b² = 275 ≡ 16 (mod 37) ≠ 0"],
        ["Gerador G", "(4, 1)"],
        ["Ordem de G", "20"],
        ["Pontos afins", f"{len(points)} + ponto no infinito O  →  |E| = {len(points)+1}"],
        ["Aviso", "Curva minúscula. Quebrável em segundos. Apenas didática."],
    ],
    headers=["Campo", "Valor"],
)

h("1. Pergunta")
para(
    "Como transformar as 18 mensagens-núcleo do plano 365 em pontos de uma curva elíptica, "
    "de modo que qualquer leitor da Almeida possa recalcular o ponto à mão e conferir o resultado "
    "— e, no mesmo objeto, ver derivada, integral e desvio padrão no nível de matemática "
    "elementar de vestibular?"
)

h("2. Evidência (o que é ECC)")
para(
    "A criptografia de curva elíptica (ECC) usa o grupo abeliano dos pontos de uma curva "
    "y² = x³ + ax + b sobre um corpo finito. A operação é a adição de pontos: a reta que liga "
    "P e Q corta a curva num terceiro ponto, refletido no eixo x. A segurança real (não deste "
    "brinquedo) vem do problema do logaritmo discreto elíptico: dado G e kG, achar k."
)
para(
    "Chaves menores: cerca de 256 bits em ECC equivalem a cerca de 3072 bits em RSA para "
    "128 bits de segurança (ordem de grandeza da Wikipédia). Aqui p = 37 de propósito, para caber no caderno."
)

h("3. Decisão — Curva Almeida")
para("Equação escolhida:  y² ≡ x³ + 2x + 3  (mod 37).")
para(
    "Sobre os reais, o mesmo polinômio fatora: x³ + 2x + 3 = (x + 1)(x² − x + 3). "
    "O trinômio tem discriminante 1 − 12 < 0, sempre positivo. Logo o ramo real existe para x ≥ −1. "
    "Isso permite derivar e integrar no mesmo objeto geométrico que o grupo finito."
)
para(
    "Gerador: G = (4, 1), ordem 20. O grupo total tem 40 pontos (39 afins + O). "
    "Trabalhamos no subgrupo cíclico gerado por G."
)

h("4. Protocolo de cifragem (mão)")
para("PASSO 1 — Leia a palavra-chave na tabela de conceitos (ou o título do dia na Almeida).")
para("PASSO 2 — Converta A=1 … Z=26 (ignore acentos, espaços e hífens). Some. Isso é a gematria latina.")
para("PASSO 3 — k = (soma + índice) mod 20. Se o resto for 0, use k = 20.")
para("PASSO 4 — Calcule o ponto C = k · G por duplicação e adição (método binário). Confira na tabela de múltiplos.")
para(
    "PASSO 5 — O par (k, C) é o selo da mensagem. Quem leu o texto reconstitui k e verifica C. "
    "Quem só vê C enfrenta o logaritmo discreto (aqui, só 20 possibilidades — por isso é brinquedo)."
)

h("5. Formato elíptico — adição e duplicação")
para("Seja P = (x1, y1), Q = (x2, y2), ambos diferentes de O, Q diferente de −P.")
para("Inclinação (adição, P ≠ Q):  λ = (y2 − y1) · (x2 − x1)^(−1)  (mod 37)")
para("Inclinação (duplicação, P = Q, y1 ≠ 0):  λ = (3x1² + a) · (2y1)^(−1)  (mod 37), com a = 2.")
para("Ponto soma R = (x3, y3):  x3 = λ² − x1 − x2,   y3 = λ(x1 − x3) − y1  (mod 37).")
para("Inverso: −P = (x1, −y1 mod 37). Identidade: O (infinito).")

para("Cálculo trabalhado — 2G, com G = (4, 1):", bold=True)
para(
    "λ = (3·4² + 2) · (2·1)^(−1) = 50 · 2^(−1). "
    "50 ≡ 13 (mod 37).  2^(−1) ≡ 19 porque 2·19 = 38 ≡ 1.  "
    "Logo λ ≡ 13·19 = 247 ≡ 25 (mod 37)."
)
para("x3 = 25² − 4 − 4 = 625 − 8 = 617 ≡ 25 (mod 37), pois 37·16 = 592 e 617 − 592 = 25.")
para("y3 = 25·(4 − 25) − 1 = 25·(−21) − 1 = −526 ≡ 29 (mod 37).")
para("Portanto 2G = (25, 29).  3G = 2G + G. Confira na tabela.")

rows_mult = []
for mlt in multiples:
    pt = mlt["point"]
    rows_mult.append([str(mlt["k"]), "O (infinito)" if pt is None else f"({pt[0]}, {pt[1]})"])
kv_table(rows_mult, headers=["k", "k · G = (x, y)  em F37"])

h("6. Mensagens principais cifradas")
para(
    "Dezoito conceitos-núcleo do plano 365. Cada linha é recalculável: some as letras da chave, "
    "some o índice, reduza módulo 20, leia kG."
)

rows_msg = []
for m in encoded:
    pt = m["point"]
    rows_msg.append(
        [
            str(m["id"]),
            m["key"],
            str(m["sum"]),
            str(m["k"]),
            "O" if pt is None else f"({pt[0]}, {pt[1]})",
            m["idea"],
        ]
    )
kv_table(rows_msg, headers=["#", "Chave", "Soma", "k", "Ponto C = kG", "Ideia central"])

para("Detalhe da soma (exemplo 1 — DEUS): D=4, E=5, U=21, S=19 → soma = 49. k = (49 + 1) mod 20 = 10. C = 10G = (24, 0).")
para(
    "Detalhe da soma (exemplo 12 — CRUZ): C=3, R=18, U=21, Z=26 → soma = 68. "
    "k = (68 + 12) mod 20 = 20. C = 20G = O, o infinito — a cruz como identidade do grupo: "
    "o ponto que some e devolve o mesmo ponto. Símbolo proposital."
)

h("7. Derivada, integral e desvio padrão")
para("Nível vestibular. A curva real é a mesma equação, agora em R.")

para("7.1 Derivada implícita", bold=True)
para("y² = x³ + 2x + 3")
para("2y y' = 3x² + 2    ⇒    y' = (3x² + 2) / (2y),  y ≠ 0.")
y_ex = st["y_real_at_4"]
dydx_ex = st["dydx_at_4"]
para(f"No ponto real de abscissa x = 4 (a mesma de G): y = √(64 + 8 + 3) = √75 = 5√3 ≈ {y_ex:.6f}.")
para(f"y'(4) = (3·16 + 2) / (2·5√3) = 50 / (10√3) = 5/√3 = (5√3)/3 ≈ {dydx_ex:.6f}.")
para(
    "Interpretação: a reta tangente à curva real tem a mesma fórmula de inclinação que a duplicação "
    "usa no corpo finito — é o mesmo λ, em dois mundos. Os números não coincidem porque um vive em R e o outro em F37."
)

para("7.2 Integral definida (área sob o ramo superior)", bold=True)
para("I = integral de −1 até 4 de √(x³ + 2x + 3) dx")
para(
    "Essa é uma integral elíptica (daí o nome histórico da curva). Não se pede forma elementar no vestibular; "
    "pede-se sentido de área e, se muito, valor numérico."
)
para(f"Valor numérico (trapézios, 2000 intervalos): I ≈ {st['integral_-1_to_4']:.6f}.")
para(f"Comprimento de arco do mesmo intervalo: L ≈ {st['arc_-1_to_4']:.6f}.")

para("7.3 Desvio padrão dos pontos finitos", bold=True)
para("Seja {Pi = (xi, yi)} o conjunto dos 39 pontos afins de E(F37).")
para("μx = (1/n) Σ xi     σx = √[(1/n) Σ (xi − μx)²]     (desvio populacional)")
para(f"n = 39    μx = {st['mu_x']:.6f}    σx = {st['sd_x']:.6f}")
para(f"μy = {st['mu_y']:.6f}    σy = {st['sd_y']:.6f}")
para(
    f"Nos 18 escalares k das mensagens: μk = {st['mu_k']:.4f}    σk = {st['sd_k']:.4f}    "
    f"Σk = {st['sum_of_k']}    Σ gematrias = {st['sum_of_key_gematria']}"
)
para(
    "O desvio padrão mede o espalhamento das abscissas no corpo. "
    "Numa curva de produção (p ~ 2^256) esse espalhamento é quase uniforme; aqui você vê o grão."
)

h("8. Como decifrar pela leitura da Bíblia")
para("1. Abra a Almeida na referência da mensagem (coluna da ideia central).")
para("2. Reconheça a palavra-chave (Deus, Cruz, Graça…).")
para("3. Some A=1…Z=26. Some o índice da tabela. Reduza módulo 20.")
para("4. Localize kG na tabela de múltiplos. Se coincidir com a coluna Ponto, a cifra está íntegra.")
para("5. Dias do cronograma 365 usam a mesma regra sobre o TÍTULO DO DIA. Exemplo Dia 001 — NO PRINCÍPIO:")

d1 = day_enc[0]
pt1 = d1["point"]
para(f"Dia {d1['day']:03d} | {d1['title']} | leitura: {d1['reading']}")
para(
    f"Soma do título = {d1['sum']}    k = (soma + 1) mod 20 = {d1['k']}    "
    f"ponto = {tuple(pt1) if pt1 else 'O'}"
)
if d1.get("wisdom"):
    para("Palavra de sabedoria: " + d1["wisdom"])

h("9. Imagem enigmática")
para(
    "A constelação abaixo é a soma aritmética das 18 chaves, projetada sobre o ramo real. "
    "Cada estrela é um kG. A polilinha segue a ordem do enredo bíblico. Os números 1–18 são o único alfabeto. "
    "Quem leu o plano reconhece o caminho; quem não leu vê só uma curva. O símbolo ∞ marca a Cruz (20G = O)."
)
doc.add_picture(str(img_doc_path), width=Cm(17.2))
p = doc.add_paragraph()
r = p.add_run(
    "Figura. Constelação Sem Jaula — 18 pontos kG sobre y² = x³ + 2x + 3. "
    "σx dos 39 pontos afins impresso na legenda."
)
set_run(r, size=9, italic=True, color=MUTED)

h("10. Síntese prática — o que fazer amanhã")
para("1. Abra o Dia 001 na Almeida (Gênesis 1:1-5; João 1:1-5).")
para("2. Recalcule à mão: DEUS → 49; k = 10; 10G = (24, 0) na tabela.")
para(
    "3. No simulador de probabilidade, lance uma moeda 10, 100 e 1000 vezes. Compare a frequência com 1/2. "
    "A intuição de ECC (kG parece aleatório) e a intuição de frequência (N grande abraça o teórico) "
    "são a mesma disciplina: o acaso visível, a estrutura invisível."
)
para("4. Assine o manifesto do Dia 365 só depois de verificar pelo menos 3 pontos da tabela.")

h("11. Erros a evitar")
para("VERSÍCULO SOLTO — cifrar uma frase fora do capítulo.")
para("CURVA DE PRODUÇÃO — não reutilize p = 37, G = (4,1) para proteger dado real.")
para("CONFUNDIR R COM F_p — a derivada vive nos reais; a cifra vive no módulo 37. O λ é análogo, os números não são os mesmos.")
para("PERFECCIONISMO DE SEQUÊNCIA — perdeu um dia, continue do marcador.")

h("12. Catálogo")
kv_table(
    [
        ["Produto", "Cifra Elíptica 365 + Simulador de Intuição"],
        ["Versão", "1.0 — 14/09/2026"],
        ["Padrão", "TURBO TEMI 360 (pergunta → evidência → decisão → protocolo → cálculo → catálogo)"],
        ["Tradução", "João Ferreira de Almeida (referências; texto não reproduzido)"],
        ["Matemática", "Weierstrass, grupo abeliano, derivada implícita, integral definida, desvio populacional"],
        ["Público", "Leitor do plano 365; vestibular (função, derivada, integral, estatística descritiva)"],
        ["Licença de uso", "Didático. ECC real exige curvas padronizadas (P-256, secp256k1) e bibliotecas auditadas."],
    ],
    headers=["Item", "Registro"],
)

p = doc.add_paragraph()
r = p.add_run("FIM DO PROTOCOLO. O ESFORÇO CONTINUA; A CULPA NÃO PRECISA CONTINUAR.")
set_run(r, size=10, bold=True, color=SAGE)

h("Anexo — primeiros 21 dias do cronograma, já selados")
rows_d = []
for d in day_enc[:21]:
    pt = d["point"]
    rows_d.append(
        [
            f"{d['day']:03d}",
            d["title"][:42],
            str(d["sum"]),
            str(d["k"]),
            "O" if pt is None else f"({pt[0]}, {pt[1]})",
        ]
    )
kv_table(rows_d, headers=["Dia", "Título", "Soma", "k", "kG"])

out_docx = DL / "Biblia_365_Cifra_Eliptica_TEMI360.docx"
doc.save(out_docx)
print("docx", out_docx, out_docx.stat().st_size)
